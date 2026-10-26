"""Plain-language results guide: thesis_outputs/THESIS_GUIDE.docx.

A short, easy-to-read Word document for the student and supervisor: the study in
one page, each result with the figure/table to cite, a figure-reading guide, a
glossary, limitations and open decisions. Every number is read from the result
files written by the earlier stages; only the wording around them is fixed, and
sentences that interpret a result depend on the numbers (e.g. significance).
"""
from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd

from . import catalog
from .experiments import MODEL_LABELS
from .feature_selection import FS_LABELS
from .utils import Paths, get_logger, git_commit

ACCENT = "1F4E79"
CALLOUT_FILL = "EAF1F8"
HEADER_FILL = "D9E2F3"

# What to look for in each catalogued figure (wording only; no results).
LOOK_FOR = {
    "methodology_flowchart": "The whole study on one page: data → curation → statistics and machine learning "
                             "→ explanation → web app. Use it to introduce Chapter 3.",
    "curation_flow": "How many responses were removed at each cleaning step and why.",
    "bdi_mapping": "How a BDI total score becomes 'depressed / not depressed' and one of six severity classes.",
    "preprocessing": "Which steps are done once per participant (scoring) and which are re-learned inside every "
                     "training fold (encoding, scaling, SMOTE, feature selection) so no test information leaks.",
    "hybrid_framework": "The proposed model: a feature-selection step chooses the inputs, four base learners are "
                        "tuned, and their predictions are combined by voting or stacking.",
    "nested_cv": "Two loops: the inner loop tunes settings, the outer loop measures performance on data the "
                 "model never saw.",
    "web_architecture": "How the Streamlit app, the saved models and the questionnaire fit together.",
    "class_distribution": "How unequal the classes are. Few 'not depressed' participants means that class is "
                          "hard to learn and its results are less certain.",
    "score_distribution": "The spread of BDI and FCRI totals; dashed lines are the severity cut-offs.",
    "fcri_by_severity": "Whether fear of recurrence rises from the 'Normal' to the 'Extreme' class. A steady rise "
                        "supports the link between FCR and depression.",
    "correlation_heatmap": "Darker cells = stronger correlation. Look at the BDI row: which FCRI subscales move "
                           "most with depression?",
    "fcri_bdi_scatter": "Each dot is a participant. A tight upward cloud means high FCR goes with high BDI.",
    "logistic_forest": "Odds ratios to the right of 1 raise the odds of depression, to the left lower them. "
                       "Lines crossing 1 are not significant. Very wide lines mean too little data.",
    "featureset_comparison": "The key leakage figure. Set A (with BDI items) is high because the label is made "
                             "from those items. Compare B, C and D to see what FCR and patient details add.",
    "model_fs_comparison": "Each cell is a model + feature-selection method; darker = better. Find the best "
                           "combination (the hybrid).",
    "roc_curves": "Curves closer to the top-left corner separate depressed from not depressed better "
                  "(higher AUC).",
    "pr_curves": "Precision-recall for the rare 'not depressed' class; compare each curve with the dashed "
                 "chance line.",
    "confusion_binary": "Rows are the true class, columns the prediction. Off-diagonal cells are mistakes.",
    "confusion_severity": "Most mistakes should sit next to the diagonal (one class away), which is less "
                          "serious than a big jump.",
    "cd_diagram": "Models further left rank better. Models joined by a thick bar are not significantly "
                  "different.",
    "calibration": "If the curve follows the diagonal, a predicted probability of 0.8 really means about 80% "
                   "are depressed.",
    "learning_curve": "If the validation line is still rising at the right, more participants would help. "
                      "A big gap to the training line means over-fitting.",
    "error_analysis": "Errors are most common for participants whose BDI score lies close to a cut-off, i.e. "
                      "the borderline cases.",
    "imbalance_ablation": "Bars show how much macro F1 changes when something is removed or changed. "
                          "Only bars marked significant matter.",
    "fs_stability": "How often each feature was chosen across training folds; 1.0 = chosen every time.",
    "shap_beeswarm": "Each dot is a participant. Dots to the right push the prediction towards depression; "
                     "the colour shows whether the answer on that item was high or low.",
    "shap_bar": "The average importance of each feature for the final model.",
    "shap_dependence": "How the effect of one item changes as the answer goes from 0 to 4.",
    "shap_waterfall": "Why the model made one specific prediction: each bar pushes the probability up or down.",
    "shap_groups": "Share of the model's reasoning that comes from BDI, FCRI, demographic and clinical "
                   "variables.",
}

GLOSSARY = [
    ("BDI", "Beck Depression Inventory: 21 questions, each scored 0–3. The total (0–63) defines both targets."),
    ("FCRI", "Fear of Cancer Recurrence Inventory: 42 questions, each scored 0–4, in seven subscales."),
    ("FCR", "Fear of cancer recurrence: worry that the cancer will come back or spread."),
    ("Target / label", "What the model tries to predict: depressed yes/no (binary) or one of six severity classes."),
    ("Feature", "An input given to the model, e.g. one FCRI answer, age group or cancer type."),
    ("Leakage", "When the model is given information that already contains the answer. Here: BDI items, "
                "because the label is calculated from them."),
    ("Feature selection (FS)", "A step that keeps only the most useful inputs (methods: mutual information, χ², "
                               "RFE, L1, Boruta)."),
    ("Ensemble", "Several models combined. Soft voting averages their probabilities; stacking trains a small "
                 "logistic-regression model on top of their predictions."),
    ("Hybrid model", "The best feature-selection method + the best ensemble, chosen by macro F1 on feature set B."),
    ("Cross-validation (CV)", "Split the data into 5 parts, train on 4 and test on 1, rotate, and repeat with "
                              "different splits."),
    ("Nested CV", "CV inside CV: the inner loop tunes the model's settings, the outer loop gives an honest score."),
    ("Class weights / SMOTE", "Two ways to stop the model ignoring small classes: give their errors more weight, "
                              "or create synthetic examples of them."),
    ("Macro F1", "Main score (0–1). F1 balances precision and recall; 'macro' averages it over classes so the "
                 "small classes count equally."),
    ("Sensitivity (recall)", "Share of truly depressed participants the model finds."),
    ("Specificity", "Share of truly not-depressed participants the model correctly clears."),
    ("ROC-AUC", "Chance that the model ranks a random depressed participant above a random non-depressed one "
                "(0.5 = guessing, 1 = perfect)."),
    ("MCC", "Matthews correlation coefficient, −1 to 1 (0 = guessing); fair under class imbalance."),
    ("Quadratic weighted κ (QWK)", "Agreement for ordered classes; a near miss is penalised less than a big miss."),
    ("Brier score", "Average squared error of the predicted probabilities; lower is better."),
    ("Mean ± SD", "Average over the outer test folds ± how much it varied between folds."),
    ("95% CI", "Range that very likely contains the true value (from bootstrap resampling of participants)."),
    ("p-value / Holm", "p < 0.05 = unlikely to be chance. Holm adjusts p-values when many tests are run at once."),
    ("Spearman ρ", "Rank correlation from −1 to 1; above 0.7 is usually called strong."),
    ("Cronbach's α", "Internal consistency of a questionnaire; above 0.7 is acceptable, above 0.9 excellent."),
    ("Ablation", "Remove one thing (e.g. demographic variables) and see whether performance drops."),
    ("SHAP", "A method that shows how much each input pushed a prediction up or down. It explains the model, "
             "not cause and effect."),
    ("Nogueira stability", "0–1; how similar the selected features are across training folds."),
]


# ----------------------------------------------------------------------------
# Small helpers
# ----------------------------------------------------------------------------
def _load(p: Path) -> Any:
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else None


def _csv(p: Path) -> pd.DataFrame | None:
    return pd.read_csv(p, encoding="utf-8-sig") if p.exists() else None


def _f(x: float, nd: int = 2) -> str:
    return f"{x:.{nd}f}"


def _pct(x: float) -> str:
    return f"{100 * x:.0f}%"


def _p(p: float) -> str:
    return "p < 0.001" if p < 0.001 else f"p = {p:.3f}"


def _first_float(s: Any) -> float:
    m = re.match(r"\s*(-?\d+(\.\d+)?)", str(s))
    return float(m.group(1)) if m else np.nan


def _fig(key: str, cfg: dict) -> str:
    return f"Figure {catalog.figure(key, cfg)['number']}"


def _tab(key: str, cfg: dict) -> str:
    return f"Table {catalog.table(key, cfg)['number']}"


def _comparisons(kr: dict, sub: str, task: str) -> dict[str, dict]:
    """Hybrid vs the same ensemble without FS ('no_fs'), and set B vs A/C/D (same ensemble, no FS)."""
    out = {}
    for r in kr.get("significance", []):
        if r["subset"] != sub or r["task"] != task:
            continue
        ra, rb = r["a"].rsplit(":", 1)[0], r["b"].rsplit(":", 1)[0]
        if ra == rb:
            continue
        if ra.startswith("fs__") and rb.startswith("featuresets__"):
            out["no_fs"] = r
        elif ra.startswith("featuresets__"):
            out[rb.split("__")[3]] = r
    return out


# ----------------------------------------------------------------------------
# Word helpers
# ----------------------------------------------------------------------------
def _shade(cell, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


class Doc:
    """Thin wrapper around python-docx with the guide's house style."""

    def __init__(self) -> None:
        from docx import Document
        from docx.shared import Cm, Pt, RGBColor
        self.Pt, self.Cm, self.RGB = Pt, Cm, RGBColor
        self.d = Document()
        st = self.d.styles["Normal"]
        st.font.name = "Calibri"
        st.font.size = Pt(11)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.line_spacing = 1.15
        for lvl, size in ((1, 16), (2, 13), (3, 11.5)):
            h = self.d.styles[f"Heading {lvl}"]
            h.font.name = "Calibri"
            h.font.size = Pt(size)
            h.font.color.rgb = RGBColor.from_string(ACCENT)
        sec = self.d.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.3)
        sec.top_margin = sec.bottom_margin = Cm(2.0)

    def h(self, text: str, level: int = 1, new_page: bool = False) -> None:
        self.d.add_heading(text, level=level).paragraph_format.page_break_before = new_page

    def p(self, text: str = "", bold_lead: str | None = None, italic: bool = False, size: float | None = None):
        par = self.d.add_paragraph()
        if bold_lead:
            par.add_run(bold_lead).bold = True
        r = par.add_run(text)
        r.italic = italic
        if size:
            r.font.size = self.Pt(size)
        return par

    def bullets(self, items: list[str | tuple[str, str]], numbered: bool = False) -> None:
        for i, it in enumerate(items, 1):
            if numbered:  # manual numbers: python-docx "List Number" continues across lists
                par = self.d.add_paragraph()
                par.paragraph_format.left_indent = self.Cm(0.9)
                par.paragraph_format.first_line_indent = self.Cm(-0.6)
                par.add_run(f"{i}.	")
            else:
                par = self.d.add_paragraph(style="List Bullet")
            par.paragraph_format.space_after = self.Pt(3)
            if isinstance(it, tuple):
                par.add_run(it[0]).bold = True
                par.add_run(it[1])
            else:
                par.add_run(it)

    def callout(self, title: str, lines: list[str]) -> None:
        t = self.d.add_table(rows=1, cols=1)
        c = t.rows[0].cells[0]
        _shade(c, CALLOUT_FILL)
        c.paragraphs[0].add_run(title).bold = True
        for ln in lines:
            par = c.add_paragraph(style="List Bullet")
            par.add_run(ln)
            par.paragraph_format.space_after = self.Pt(2)
        self._spacer()

    def table(self, header: list[str], rows: list[list[Any]], widths_cm: list[float] | None = None,
              font: float = 10) -> None:
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        t = self.d.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        for j, hd in enumerate(header):
            c = t.rows[0].cells[j]
            _shade(c, HEADER_FILL)
            r = c.paragraphs[0].add_run(str(hd))
            r.bold = True
            r.font.size = self.Pt(font)
        for row in rows:
            cells = t.add_row().cells
            for j, v in enumerate(row):
                par = cells[j].paragraphs[0]
                r = par.add_run("" if v is None else str(v))
                r.font.size = self.Pt(font)
                if j > 0 and re.match(r"^[\d.\-–+−±% <p=()*]+$", str(v) or "x"):
                    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for row in t.rows:
            for j, c in enumerate(row.cells):
                for par in c.paragraphs:
                    par.paragraph_format.space_after = self.Pt(1)
                if widths_cm:
                    c.width = self.Cm(widths_cm[j])
        if len(t.rows) <= 16:  # keep short tables on one page
            for row in t.rows[:-1]:
                for c in row.cells:
                    for par in c.paragraphs:
                        par.paragraph_format.keep_with_next = True
        self._spacer()

    def _spacer(self) -> None:
        par = self.d.add_paragraph()
        par.paragraph_format.space_after = self.Pt(4)
        par.paragraph_format.line_spacing = 1.0
        par.add_run("").font.size = self.Pt(2)

    def figure(self, png: Path, label: str, caption: str, max_w_cm: float = 15.5, max_h_cm: float = 19.0) -> None:
        if not png.exists():
            self.p(f"[{label} not found: {png.name} — run the pipeline]", italic=True)
            return
        from PIL import Image
        with Image.open(png) as im:
            w, h = im.size
        width = min(max_w_cm, max_h_cm * w / h)
        self.d.add_picture(str(png), width=self.Cm(width))
        self.d.paragraphs[-1].alignment = 1
        par = self.d.add_paragraph()
        par.alignment = 1
        r = par.add_run(f"{label}. ")
        r.bold = True
        r.font.size = self.Pt(9.5)
        r2 = par.add_run(caption)
        r2.italic = True
        r2.font.size = self.Pt(9.5)


# ----------------------------------------------------------------------------
# Content
# ----------------------------------------------------------------------------
def build(ds, cfg: dict[str, Any], paths: Paths) -> Path:
    """Write thesis_outputs/THESIS_GUIDE.docx from the result files."""
    log = get_logger()
    R, T = paths.results, paths.tables
    kr = _load(R / "key_results.json")
    counts = _load(R / "class_counts.json")
    stats = _load(R / "statistics_summary.json")
    quality = (_load(R / "data_quality.json") or {}).get("quality")
    shap = _load(R / "shap_summary.json")
    if not (kr and counts and stats):
        raise FileNotFoundError("Result files missing. Run:  python -m thesis_analysis.run_all --stage all")

    subs = cfg["run_subsets"]
    prim = cfg["primary_subset"]
    sec = next((s for s in subs if s != prim), None)
    alpha = cfg["cv"]["alpha"]
    key = cfg["key_feature_set"]
    cv = cfg["cv"]
    n_folds = cv["outer_splits"] * cv["outer_repeats"]
    lab = {s: catalog.subset_label(cfg, s) for s in subs}
    cut = cfg["targets"]["binary_cutoff"]

    def hyb(sub, task):
        return kr[sub][task]["hybrid"]

    def hm(sub, task, metric):
        return hyb(sub, task)["metrics"][metric]["mean"]

    def hname(sub, task):
        h = hyb(sub, task)
        return f"{FS_LABELS[h['fs']]} + {MODEL_LABELS[h['model']]}"

    def full_item(label: str) -> str:
        m = re.match(r"^FCRI (\d+):", label)
        col = f"FCRI_{m.group(1)}" if m else None
        if col and col in getattr(ds, "item_text", {}):
            return f"FCRI {m.group(1)}: {ds.item_text[col]}"
        return label

    def fig_png(k: str) -> Path:
        return paths.figures / f"{catalog.figure(k, cfg)['name']}.png"

    def cap(k: str) -> str:
        return catalog.figure(k, cfg)["caption"]

    D = Doc()
    d = D.d

    # ---------------- Title ----------------
    t = d.add_paragraph()
    r = t.add_run("Thesis Results Guide")
    r.bold = True
    r.font.size = D.Pt(24)
    r.font.color.rgb = D.RGB.from_string(ACCENT)
    D.p("Hybrid Machine Learning Framework for Detection of Depression and its Severity in Cancer Patients "
        "using Fear of Cancer Recurrence (FCR)", size=13)
    D.p("MS Bioinformatics thesis, NUST. A plain-language companion to the figures and tables in "
        "thesis_outputs/.", italic=True)
    D.p(f"Generated {datetime.now():%d %B %Y} from git commit {git_commit()}. Every number in this document is "
        "read from the analysis result files; to update it, re-run  python -m thesis_analysis.run_all --stage "
        "report.", italic=True, size=9.5)
    D.callout("How to use this guide", [
        "Part 1 tells the whole story on one page — read it first.",
        "Parts 2–5 explain each result in simple words and name the figure or table to cite in the thesis.",
        "Part 6 explains how to read every figure; Part 7 is a glossary of the technical terms.",
        "Parts 8–9 list the limitations and the decisions to agree with your supervisor before submission.",
        f"Unless stated otherwise, results are for '{lab[prim]}' (n = {counts[prim]['n']}), the primary analysis.",
    ])

    # ---------------- Part 1 ----------------
    D.h("1. The study in one page")
    D.p("Can fear of cancer recurrence (FCR), together with basic demographic and clinical information, "
        "detect depression and estimate its severity in people with cancer?", bold_lead="Research question. ")
    D.p(f"{counts[prim]['n']} participants answered the BDI depression questionnaire and the FCRI "
        "fear-of-recurrence questionnaire, plus questions on gender, age, province, cancer type, time since "
        f"diagnosis, cancer status, treatment and medication. A BDI total of {cut} or more was counted as "
        "depressed; the BDI total was also split into six severity classes.", bold_lead="Data. ")
    D.p("(1) cleaned and scored the data; (2) tested statistically how FCR relates to depression; "
        "(3) trained machine-learning models with and without the depression questions, adding a "
        "feature-selection step and combining models into ensembles (the 'hybrid' framework); "
        "(4) measured performance honestly with repeated nested cross-validation and explained the final "
        "models with SHAP.", bold_lead="What was done. ")

    st = stats[prim]
    rho = st["spearman_fcri_total_bdi"]["rho"]
    cmpb = _comparisons(kr, prim, "binary")
    cmps = _comparisons(kr, prim, "severity6")
    ans = [
        f"Fear of recurrence and depression rise together strongly (Spearman ρ = {_f(rho)}).",
        f"Without seeing any depression questions, the hybrid model ({hname(prim, 'binary')}) detected "
        f"depression with macro F1 {_f(hm(prim, 'binary', 'f1_macro'))}: it found "
        f"{_pct(hm(prim, 'binary', 'recall'))} of depressed participants and correctly cleared "
        f"{_pct(hm(prim, 'binary', 'specificity'))} of non-depressed participants "
        f"(ROC-AUC {_f(hm(prim, 'binary', 'roc_auc'))}).",
        f"Estimating the exact severity class is harder: macro F1 {_f(hm(prim, 'severity6', 'f1_macro'))}, "
        f"but most mistakes are near misses (quadratic weighted κ {_f(hm(prim, 'severity6', 'qwk'))}).",
    ]
    if "C" in cmpb and "D" in cmpb:
        bc, bd = cmpb["C"], cmpb["D"]
        ans.append(
            "The FCRI answers carry the signal. FCRI alone performed "
            + ("about the same as" if bc["p_corrected_t_holm"] >= alpha else "significantly differently from")
            + f" FCRI + demographic + clinical (binary macro F1 {_f(bc['mean_b'], 3)} vs {_f(bc['mean_a'], 3)}), "
            f"while demographic + clinical information alone reached only {_f(bd['mean_b'], 3)}.")
    if "no_fs" in cmpb:
        nf = cmpb["no_fs"]
        ans.append(
            "Feature selection made the model simpler but "
            + ("did not make it significantly better" if nf["p_corrected_t_holm"] >= alpha
               else "also made it significantly better")
            + f" than the same ensemble without it (difference {nf['diff']:+.3f}, {_p(nf['p_corrected_t_holm'])}).")
    if sec:
        ans.append(
            f"Caution: when the {counts[prim]['n'] - counts[sec]['n']} rows that match a mock-data generator "
            f"are removed, binary macro F1 falls to {_f(hm(sec, 'binary', 'f1_macro'))} and severity to "
            f"{_f(hm(sec, 'severity6', 'f1_macro'))}. Which rows are real must be settled before submission.")
    D.callout("The answer in brief", ans)

    # ---------------- Part 2: data ----------------
    D.h("2. The data")
    D.p("All analyses were run twice, on two versions of the dataset:")
    bl = [(f"{lab[prim]} (n = {counts[prim]['n']}). ", "Every cleaned response. This is the primary analysis.")]
    if sec and quality:
        bl.append((f"{lab[sec]} (n = {counts[sec]['n']}). ",
                   f"{quality['generator_matched_rows']} rows have a name and all 21 BDI answers identical to "
                   "the output of a mock-data generator found in the project's history, so they may not be real "
                   "participants. This version leaves them out."))
    D.bullets(bl)
    rows = [["Participants", *[counts[s]["n"] for s in subs]],
            [f"Depressed (BDI ≥ {cut})", *[counts[s]["binary"]["1"] for s in subs]],
            [f"Not depressed (BDI < {cut})", *[counts[s]["binary"]["0"] for s in subs]]]
    for name, lo, hi in cfg["targets"]["severity_bounds"]:
        rows.append([f"   {name} (BDI {lo}–{hi})", *[counts[s]["severity6"].get(name, "–") for s in subs]])
    rows.append(["Mean BDI total (SD)", *[f"{counts[s]['bdi_total_mean']:.1f} ({counts[s]['bdi_total_sd']:.1f})"
                                           for s in subs]])
    rows.append(["Mean FCRI total (SD)", *[f"{counts[s]['fcri_total_mean']:.1f} ({counts[s]['fcri_total_sd']:.1f})"
                                            for s in subs]])
    D.table(["", *[lab[s] for s in subs]], rows, [7.4, 4.4, 4.4])
    small = int(counts[prim]["binary"]["0"])
    D.p(f"Most participants are depressed; only {small} are not. Classes this unequal make the "
        "'not depressed' group hard to learn, which is why the main score is macro F1 (it gives each class "
        f"equal weight) and why class weights are used. See {_fig('class_distribution', cfg)} and "
        f"{_tab('characteristics_' + prim, cfg)}.", bold_lead="What this means. ")
    rel = _csv(T / f"{catalog.table('reliability', cfg)['name']}.csv")
    if rel is not None:
        rp = rel[rel.Subset == lab[prim]]
        bdi_a = rp[rp.Scale.str.startswith("BDI")]["Cronbach's alpha"]
        fc_a = rp[rp.Scale.str.startswith("FCRI total")]["Cronbach's alpha"]
        if len(bdi_a) and len(fc_a):
            D.p(f"Both questionnaires are internally consistent (Cronbach's α: BDI {_f(float(bdi_a.iloc[0]))}, "
                f"FCRI {_f(float(fc_a.iloc[0]))}; {_tab('reliability', cfg)}).", bold_lead="Reliability. ")
    if quality:
        D.h("Data problems found", 2)
        D.bullets([
            f"{quality['bdi_items_6_7_rescored']} answers to BDI items 6 and 7 had been scored wrongly by the "
            "original parser (\"I don't feel ...\" was scored as a symptom). They were corrected here; the saved "
            "app models in models/ were trained before this fix.",
            f"{quality['male_with_female_specific_cancer']} men report ovarian or cervical cancer and "
            f"{quality['female_with_prostate_cancer']} women report prostate cancer.",
            f"{quality['under_18']} participants are under 18 years old.",
            f"{quality['bdi_straightlining']} participants gave the same answer to all 21 BDI questions.",
            f"Duplicate responses found: {quality['exact_duplicate_rows']}. "
            f"{quality['participants_with_any_missing_fcri']} participants had a few missing FCRI answers, "
            "which were filled with their own subscale average.",
        ])
        D.p(f"These are listed in {_tab('data_quality', cfg)}. They were kept in the analysis but should be "
            "discussed with your supervisor (Part 9).", italic=True)

    # ---------------- Part 3: methods ----------------
    D.h("3. What was done, in plain words", new_page=True)
    D.figure(fig_png("methodology_flowchart"), _fig("methodology_flowchart", cfg), cap("methodology_flowchart"),
             max_h_cm=16)
    D.h("Step 1 — Two targets from one score", 2)
    bands = ", ".join(f"{a} {lo}–{hi}" for a, lo, hi in cfg["targets"]["severity_bounds"])
    D.p(f"The BDI total decides both things the model must predict: depressed if BDI ≥ {cut} (binary), and "
        f"a severity class ({bands}). See {_fig('bdi_mapping', cfg)}.")
    D.h("Step 2 — Four feature sets, to avoid fooling ourselves", 2)
    D.p("Because the labels are calculated from the BDI answers, a model that sees those answers simply "
        "adds them up again; this is called leakage. Four sets of inputs were therefore compared:")
    D.table(["Set", "Inputs", "Why it is there"], [
        ["A", "BDI + FCRI + demographic + clinical", "Upper bound only — contains the answer (leakage)."],
        ["B", "FCRI + demographic + clinical", "The honest test of the thesis claim. Main set."],
        ["C", "FCRI only", "Does fear of recurrence alone carry the signal?"],
        ["D", "Demographic + clinical", "What can be done without any questionnaire?"],
    ], [1.3, 6.0, 9.0])
    D.h("Step 3 — The hybrid framework", 2)
    D.p("For every training split the pipeline (1) one-hot encodes categories and scales answers to 0–1, "
        "(2) chooses the most useful inputs with one of five feature-selection methods (mutual information, χ², "
        "RFE, L1 logistic regression, Boruta), (3) tunes four learners (logistic regression, random forest, "
        "XGBoost, SVM) and (4) combines them by soft voting or stacking. The best feature-selection method + "
        f"best ensemble on set {key} is the hybrid model. See {_fig('hybrid_framework', cfg)} and "
        f"{_tab('search_space', cfg)}.")
    D.h("Step 4 — Honest testing with nested cross-validation", 2)
    D.p(f"The data were split into {cv['outer_splits']} parts; each part in turn was the test set and the rest "
        f"the training set, and this was repeated {cv['outer_repeats']} times with different splits "
        f"({n_folds} test results per model). Inside each training set a second {cv['inner_splits']}-fold loop "
        "tuned the settings, so the test part was never used for any decision. Every step that learns from "
        "data (scaling, SMOTE, feature selection) was re-done inside each training split. "
        f"See {_fig('nested_cv', cfg)} and {_fig('preprocessing', cfg)}. A separate hold-out test set was not "
        "used because the 'not depressed' group is too small to split off reliably.")
    D.h("Step 5 — Explaining the model", 2)
    D.p("The final hybrid models were explained with SHAP, which shows how much each answer pushed each "
        "prediction towards or away from depression.")

    # ---------------- Part 4: results ----------------
    D.h("4. The results, one at a time", new_page=True)

    # 4.1 statistics
    D.h("4.1 Fear of recurrence is strongly linked with depression", 2)
    lines = [f"Spearman correlation between FCRI total and BDI total: ρ = {_f(rho)} "
             f"({_p(st['spearman_fcri_total_bdi']['p'])}).",
             f"{st['n_significant_bivariate_holm']} of {st['n_bivariate_tests']} comparisons between depressed and "
             "non-depressed participants were significant after Holm correction."]
    epv = st["logistic"].get("events_per_variable")
    if epv is not None:
        lines.append(f"The multivariable logistic regression has {epv} 'not depressed' cases per model parameter"
                     + (" — below the usual minimum of 10, so its odds ratios are unstable; report them with care."
                        if epv < 10 else "."))
    D.callout("Key numbers", lines)
    D.p(f"Cite {_fig('fcri_bdi_scatter', cfg)}, {_fig('fcri_by_severity', cfg)}, {_fig('correlation_heatmap', cfg)}, "
        f"{_tab('bivariate_binary_' + prim, cfg)} and {_tab('spearman_' + prim, cfg)}.")
    D.figure(fig_png("fcri_bdi_scatter"), _fig("fcri_bdi_scatter", cfg), cap("fcri_bdi_scatter"), max_w_cm=12)

    # 4.2 feature sets
    D.h("4.2 Which information matters? (feature sets A–D)", 2)
    rows = []
    for s in subs:
        for task, tname in (("binary", "Binary"), ("severity6", "Severity")):
            bf = kr[s][task]["best_by_feature_set"]
            rows.append([lab[s] if task == "binary" else "", tname,
                         *[_f(bf[k]["mean"]) if k in bf else "–" for k in "ABCD"]])
    D.table(["Data", "Task", "A (with BDI)", "B (FCRI + info)", "C (FCRI only)", "D (info only)"], rows,
            [4.2, 1.9, 2.4, 2.6, 2.4, 2.4], font=9.5)
    D.p(f"Best macro F1 for each feature set (mean over {n_folds} outer test folds; for set {key} this is the "
        "hybrid model).", italic=True, size=9.5)
    expl = ["Set A is high because the model can add up the BDI answers again — it is not a fair test (leakage)."]

    def diff_txt(code: str) -> str:
        parts = []
        for nm, cm in (("binary", cmpb), ("severity", cmps)):
            if code in cm:
                parts.append(f"{cm[code]['diff']:+.3f} for {nm} ({_p(cm[code]['p_corrected_t_holm'])})")
        return " and ".join(parts)

    if "C" in cmpb:
        expl.append(f"Set {key} minus set C (effect of adding demographic + clinical to FCRI): {diff_txt('C')}.")
    if "D" in cmpb:
        expl.append(f"Set {key} minus set D (effect of the FCRI answers): {diff_txt('D')}.")
    if ("C" in cmpb and cmpb["C"]["p_corrected_t_holm"] >= alpha
            and "D" in cmpb and cmpb["D"]["p_corrected_t_holm"] < alpha):
        expl.append("Plain meaning: almost all of the predictive information comes from the FCRI answers; "
                    "demographic and clinical details add little.")
    D.callout("What it shows", expl)
    D.p("These comparisons use the same ensemble without feature selection on both sides, so only the inputs "
        f"differ. Cite {_fig('featureset_comparison', cfg)} and {_tab('featureset_results', cfg)}.")
    D.figure(fig_png("featureset_comparison"), _fig("featureset_comparison", cfg), cap("featureset_comparison"))

    # 4.3 hybrid performance
    D.h("4.3 How well does the hybrid model work?", 2)
    D.p(f"Feature set {key} (no BDI answers). Values are the mean over {n_folds} outer test folds; the 95% "
        "confidence interval is in brackets.")

    def mci(sub, task, m, pct=False):
        d_ = hyb(sub, task)["metrics"][m]
        lo, hi = (d_.get("ci") or [None, None])[:2]
        fmt = _pct if pct else ((lambda x: _f(x, 3)) if m == "brier" else _f)
        ok = lo is not None and not (isinstance(lo, float) and np.isnan(lo))
        return fmt(d_["mean"]) + (f" ({fmt(lo)}–{fmt(hi)})" if ok else "")

    D.h("Binary: depressed or not", 3)
    D.table(["", *[lab[s] for s in subs]], [
        ["Hybrid model", *[hname(s, "binary") for s in subs]],
        ["Macro F1 (main score)", *[mci(s, "binary", "f1_macro") for s in subs]],
        ["Balanced accuracy", *[mci(s, "binary", "balanced_accuracy") for s in subs]],
        ["Sensitivity — depressed found", *[mci(s, "binary", "recall", True) for s in subs]],
        ["Specificity — non-depressed cleared", *[mci(s, "binary", "specificity", True) for s in subs]],
        ["ROC-AUC", *[mci(s, "binary", "roc_auc") for s in subs]],
        ["MCC", *[mci(s, "binary", "mcc") for s in subs]],
        ["Brier score (lower = better)", *[mci(s, "binary", "brier") for s in subs]],
    ], [5.6, 5.3, 5.3], font=9.5)
    D.h("Severity: six classes", 3)
    D.table(["", *[lab[s] for s in subs]], [
        ["Hybrid model", *[hname(s, "severity6") for s in subs]],
        ["Macro F1 (main score)", *[mci(s, "severity6", "f1_macro") for s in subs]],
        ["Balanced accuracy", *[mci(s, "severity6", "balanced_accuracy") for s in subs]],
        ["Quadratic weighted κ", *[mci(s, "severity6", "qwk") for s in subs]],
        ["Macro ROC-AUC", *[mci(s, "severity6", "roc_auc_ovr") for s in subs]],
    ], [5.6, 5.3, 5.3], font=9.5)

    pc = _csv(T / f"{catalog.table('per_class_' + prim, cfg)['name']}.csv")
    err = _csv(T / f"{catalog.table('error_analysis', cfg)['name']}.csv")
    lines = [f"The model finds {_pct(hm(prim, 'binary', 'recall'))} of depressed participants and correctly "
             f"clears {_pct(hm(prim, 'binary', 'specificity'))} of the non-depressed — the small group is the "
             "harder one."]
    if pc is not None and "F1" in pc:
        pc["F1v"] = pc["F1"].map(_first_float)
        best, worst = pc.loc[pc.F1v.idxmax()], pc.loc[pc.F1v.idxmin()]
        lines.append(f"Severity: easiest class '{best.Class}' (F1 {_f(best.F1v)}), hardest class "
                     f"'{worst.Class}' (F1 {_f(worst.F1v)}). The middle classes cover narrow BDI ranges, so "
                     "small differences in answers move a participant across a boundary.")
    if err is not None:
        e = err[(err.Subset == lab[prim]) & (err.Task.str.startswith("Severity")) & (err.Group == "Misclassified")]
        if len(e) and not pd.isna(e.iloc[0]["Errors to adjacent class (%)"]):
            lines.append(f"{float(e.iloc[0]['Errors to adjacent class (%)']):.0f}% of severity mistakes were only "
                         "one class away from the true class.")
        eb = err[(err.Subset == lab[prim]) & (err.Task == "Binary")].set_index("Group")
        col = "Distance to nearest cut-off (BDI points)"
        if {"Correct", "Misclassified"} <= set(eb.index):
            lines.append(f"Misclassified participants sat close to the BDI cut-off (median distance "
                         f"{eb.loc['Misclassified', col].split(' ')[0]} points vs "
                         f"{eb.loc['Correct', col].split(' ')[0]} for correctly classified ones) — the model "
                         "struggles mostly with borderline cases.")
    D.callout("What it shows", lines)
    D.p(f"Cite {_tab('binary_results_' + prim, cfg)}, {_tab('severity_results_' + prim, cfg)}, "
        f"{_tab('per_class_' + prim, cfg)}, {_fig('roc_curves', cfg)}, {_fig('confusion_binary', cfg)}, "
        f"{_fig('confusion_severity', cfg)} and {_fig('error_analysis', cfg)}.")
    D.figure(fig_png("confusion_severity"), _fig("confusion_severity", cfg), cap("confusion_severity"))

    # 4.4 comparisons
    D.h("4.4 Is the hybrid better than simpler models?", 2)
    fr = _csv(T / f"{catalog.table('friedman', cfg)['name']}.csv")
    lines = []
    if fr is not None:
        for _, r_ in fr[fr.Subset == lab[prim]].iterrows():
            lines.append(f"{r_.Task}: average ranks (1 = best) — {r_['Average ranks (1 = best)']}. "
                         f"Friedman {_p(float(r_.p))}; models whose ranks differ by less than "
                         f"{r_['Critical difference']} are not significantly different.")
    for tname, cm in (("Binary", cmpb), ("Severity", cmps)):
        if "no_fs" in cm:
            nf = cm["no_fs"]
            lines.append(f"{tname}: hybrid vs the same ensemble without feature selection: macro F1 "
                         f"{_f(nf['mean_a'], 3)} vs {_f(nf['mean_b'], 3)} ({_p(nf['p_corrected_t_holm'])}).")
    D.callout("What it shows", lines)
    no_fs_sig = any(cm.get("no_fs", {}).get("p_corrected_t_holm", 1) < alpha for cm in (cmpb, cmps))
    D.p("Plain meaning: " + (
        "the ensembles rank at the top, but feature selection did not raise the score significantly; its benefit "
        "is a smaller, more interpretable set of inputs. Describe it as 'comparable performance with fewer "
        "inputs', not as an improvement." if not no_fs_sig else
        "feature selection significantly improved at least one task — see the p-values above."))
    D.p(f"Cite {_fig('cd_diagram', cfg)}, {_tab('significance_' + prim, cfg)} and {_tab('friedman', cfg)}.")

    # 4.5 ablation
    D.h("4.5 What happens when parts are removed? (ablation)", 2)
    abl = pd.DataFrame(kr.get("ablation", []))
    if len(abl):
        D.p(f"Each row removes or changes one thing and shows the change in macro F1 compared with the hybrid "
            f"set-up ({cv['ablation_repeats']} repeats of 5-fold CV, primary data).")

        def cell(x: pd.DataFrame) -> str:
            if not len(x):
                return "–"
            r0 = x.iloc[0]
            return f"{r0.delta:+.3f}" + (" *" if r0.p_holm < alpha else "")

        rows = [[lbl_, cell(g[g.task == "binary"]), cell(g[g.task == "severity6"])]
                for lbl_, g in abl.groupby("label", sort=False)]
        D.table(["Change", "Binary Δ", "Severity Δ"], rows, [9.6, 3.2, 3.2], font=9.5)
        D.p(f"* significant after Holm correction (p < {alpha}). Negative Δ = worse without it.",
            italic=True, size=9.5)
        sig_ = abl[abl.p_holm < alpha]
        sig_labels = list(dict.fromkeys(sig_.label))
        ab_lines = [f"{len(sig_labels)} of {abl.label.nunique()} changes made a significant difference in at "
                    "least one task" + (": " + "; ".join(sig_labels) + "." if sig_labels else ".")]
        if len(sig_) and set(sig_.variant) <= {"no_fcri", "D_no_medication"}:
            ab_lines.append("Only removing the FCRI answers hurts. Removing demographic, clinical or medication "
                            "variables, the mood-overlap FCRI items, or scaling, and swapping class weights for "
                            "SMOTE, made no significant difference.")
        D.callout("What it shows", ab_lines)
        D.p(f"Cite {_fig('imbalance_ablation', cfg)}, {_tab('ablation', cfg)}, {_tab('imbalance', cfg)} "
            f"and {_tab('sensitivity', cfg)}.")
    sens = []
    for s in subs:
        k17 = kr.get(s, {}).get("sensitivity_binary17")
        k5 = kr.get(s, {}).get("sensitivity_severity5")
        if k17:
            sens.append(f"{lab[s]}: with the stricter cut-off BDI ≥ {cfg['targets']['binary_cutoff_sensitivity']}, "
                        f"best binary macro F1 = {_f(k17['mean'])}.")
        if k5:
            sens.append(f"{lab[s]}: merging the two lowest classes (five classes), best severity macro F1 = "
                        f"{_f(k5['mean'])}.")
    if sens:
        D.h("Other choices tested (sensitivity analyses)", 3)
        D.bullets(sens)

    # 4.6 SHAP
    D.h("4.6 Which questions matter most? (SHAP)", 2)
    if shap and prim in shap:
        sh = shap[prim]
        gs = sh["group_shares"]
        D.p("The five questions with the largest influence on the binary hybrid model:")
        D.bullets([full_item(k) for k, _ in list(sh["top_binary"].items())[:5]], numbered=True)
        bkey, akey = f"Binary, set {key}", "Binary, set A"
        lines = []
        if bkey in gs:
            g = gs[bkey]
            lines.append(f"In the main model (set {key}) FCRI answers account for {_pct(g.get('fcri', 0))} of the "
                         f"model's reasoning, demographic variables {_pct(g.get('demographic', 0))} and clinical "
                         f"variables {_pct(g.get('clinical', 0))}.")
            if g.get("demographic", 0) + g.get("clinical", 0) == 0:
                lines.append(f"This is because the {FS_LABELS[hyb(prim, 'binary')['fs']]} feature-selection step "
                             "kept only FCRI questions in the final model.")
        if akey in gs:
            lines.append(f"When BDI answers are allowed (set A) they take {_pct(gs[akey].get('bdi', 0))} of the "
                         "reasoning — another sign of leakage.")
        lines.append("Read the full wording of the top items above before writing the discussion: they describe "
                     "how the participant sees their risk of recurrence and how much they worry about it.")
        D.callout("What it shows", lines)
        D.p(f"Cite {_fig('shap_beeswarm', cfg)}, {_fig('shap_bar', cfg)}, {_fig('shap_dependence', cfg)}, "
            f"{_fig('shap_waterfall', cfg)}, {_fig('shap_groups', cfg)} and {_tab('key_factors_' + prim, cfg)} "
            "(which compares SHAP with the statistical tests and feature-selection frequency).")
        D.figure(fig_png("shap_bar"), _fig("shap_bar", cfg), cap("shap_bar"))

    # 4.7 calibration / learning curve
    D.h("4.7 Can the probabilities be trusted, and would more data help?", 2)
    lines = []
    cal = kr.get(prim, {}).get("binary", {}).get("calibration")
    if cal:
        lines.append(f"Brier score of the binary hybrid: {_f(cal['brier'], 3)} (0 = perfect).")
    for k, v in (kr.get("learning_curve") or {}).items():
        mid = len(v["sizes"]) // 2
        gain = v["test_mean"][-1] - v["test_mean"][mid]
        lines.append(f"{k}: macro F1 went from {_f(v['test_mean'][0])} with {v['sizes'][0]} training participants "
                     f"to {_f(v['test_mean'][mid])} with {v['sizes'][mid]} and {_f(v['test_mean'][-1])} with "
                     f"{v['sizes'][-1]} ({gain:+.3f} over the second half), so "
                     + ("the curve has levelled off and more data of the same kind would add little."
                        if gain < 0.01 else "it is still rising slowly and more data would probably help a little."))
    if lines:
        D.callout("What it shows", lines)
        D.p(f"Cite {_fig('calibration', cfg)} and {_fig('learning_curve', cfg)}.")

    # ---------------- Part 5: two subsets ----------------
    if sec:
        D.h("5. Why the two versions of the data give different answers", new_page=True)
        rows = []
        for m, nm, pct in (("f1_macro", "Binary macro F1", False), ("specificity", "Binary specificity", True),
                           ("roc_auc", "Binary ROC-AUC", False), ("mcc", "Binary MCC", False)):
            rows.append([nm, *[(_pct if pct else _f)(hm(s, "binary", m)) for s in subs]])
        rows.append(["Severity macro F1", *[_f(hm(s, "severity6", "f1_macro")) for s in subs]])
        rows.append(["Severity QWK", *[_f(hm(s, "severity6", "qwk")) for s in subs]])
        rows.append(["Not-depressed participants", *[counts[s]["binary"]["0"] for s in subs]])
        rows.append(["FCRI–BDI correlation ρ", *[_f(stats[s]["spearman_fcri_total_bdi"]["rho"]) for s in subs]])
        D.table(["", *[lab[s] for s in subs]], rows, [6.2, 5.0, 5.0])
        D.bullets([
            f"Without the generator-matched rows only {counts[sec]['binary']['0']} participants are not depressed, "
            "so the model has very few examples of that class and specificity drops sharply.",
            "The generator-matched rows may follow an artificial pattern that is easier to learn, which would make "
            "the all-rows results look better than real-world performance.",
            "Until the origin of these rows is confirmed, report both versions and state clearly which one is "
            "primary (see Part 9).",
        ])
        D.p("Appendix figures in thesis_outputs/figures/appendix/ repeat the main figures for this version.",
            italic=True)

    # ---------------- Part 6: figure guide ----------------
    D.h("6. How to read each figure", new_page=True)
    rows = [[f"Figure {m['number']}", m["caption"], LOOK_FOR.get(k, "")]
            for k, m in catalog.numbered(catalog.FIGURES, subs, "fig").items()
            if (paths.figures / f"{m['name']}.png").exists()]
    D.table(["No.", "What it is", "What to look for"], rows, [2.0, 6.2, 8.2], font=9)
    D.p("Files: thesis_outputs/figures/<name>.png (300 dpi, for Word) and .svg (vector). All are at most 16 cm "
        "wide, so they can be inserted at 100% size.", italic=True, size=9.5)

    # ---------------- Part 7: glossary ----------------
    D.h("7. Glossary")
    D.table(["Term", "Meaning in simple words"], [[a, b] for a, b in GLOSSARY], [4.0, 12.4], font=9.5)

    # ---------------- Part 8: limitations ----------------
    D.h("8. Limitations to mention in the thesis")
    lim = []
    if quality:
        lim.append(("Data origin. ", f"{quality['generator_matched_rows']} of {quality['raw_rows']} rows match "
                    "mock-generator output; the rest could not be verified from the repository alone."))
    lim += [
        ("Small 'not depressed' group. ", "Specificity and other results for this group have wide confidence "
         "intervals."),
        ("Label comes from the BDI. ", "Both targets are calculated from the BDI, so models that see BDI answers "
         "only re-compute the label; only sets B–D test real prediction."),
        ("Overlapping questions. ", "Some FCRI items (19 sadness, 26 mood) resemble depression symptoms; the "
         "ablation shows how much they matter."),
        ("Self-report, one time point, one form. ", "Both questionnaires were filled in together online, which "
         "can inflate their correlation; there is no external validation group."),
    ]
    if epv is not None and epv < 10:
        lim.append(("Regression power. ", f"Only {epv} events per parameter in the multivariable logistic "
                    "regression; its odds ratios are unstable."))
    lim += [
        ("Instrument version. ", "The BDI wording and cut-offs used are those of the first-edition BDI "
         "(0–10 normal … 41–63 extreme), not the BDI-II."),
        ("SHAP is not causation. ", "It shows what the model relies on, not what causes depression."),
    ]
    D.bullets(lim)

    # ---------------- Part 9: decisions ----------------
    D.h("9. Decisions to agree with your supervisor")
    D.bullets([
        "Which rows are genuine participants, and therefore which data version is primary in the thesis.",
        f"Whether the binary cut-off stays at BDI ≥ {cut} or moves to ≥ {cfg['targets']['binary_cutoff_sensitivity']}.",
        "Whether six severity classes are reported, or the two lowest classes are merged (five classes).",
        "Whether participants under 18 and the inconsistent gender/cancer-type answers are excluded.",
        "Whether the instrument is described as BDI-I/BDI-IA instead of BDI-II in the methodology.",
        "Whether the saved models in models/ are retrained now that the BDI items 6/7 scoring is fixed in the app parser.",
        "Whether the Streamlit app should switch to an FCRI-based hybrid model.",
    ], numbered=True)

    # ---------------- Part 10: files ----------------
    D.h("10. Where everything is")
    D.table(["What", "Where"], [
        ["Figures for the thesis", "thesis_outputs/figures/ (appendix versions in figures/appendix/)"],
        ["All tables, formatted for Word", "thesis_outputs/tables/thesis_tables.docx"],
        ["Tables as spreadsheets", "thesis_outputs/tables/*.csv"],
        ["Detailed summary + draft methodology text", "thesis_outputs/RESULTS_SUMMARY.md"],
        ["This guide", "thesis_outputs/THESIS_GUIDE.docx"],
        ["Cleaned dataset", "thesis_outputs/data/curated_dataset.csv"],
        ["Run logs (settings, versions, git commit)", "thesis_outputs/logs/"],
        ["Re-run everything", "python -m thesis_analysis.run_all --stage all"],
    ], [6.4, 10.0], font=9.5)

    out = paths.out / "THESIS_GUIDE.docx"
    d.save(out)
    log.info("Plain-language guide -> %s", out)
    return out
